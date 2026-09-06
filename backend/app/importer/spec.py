"""Импорт спецификации (Базис: CSV / XLSX / XML).

Базис выгружает вместе с DXF спецификацию: наименование детали, изделие,
габарит, кромка. Её импорт избавляет от ручного ввода изделий и карты
кромок — это половина ручной работы технолога.

Сопоставление колонок настраивается в ``config/spec_columns.yaml``;
здесь нет ни одного зашитого заголовка.
"""

from __future__ import annotations

import csv
import io
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import PurePosixPath

from app.core.config_files import load
from app.models.enums import GrainMode

_NUMERIC_FIELDS = {"length", "width", "thickness"}
_INT_FIELDS = {"qty"}


def spec_config() -> dict:
    return load("spec_columns")


@dataclass(slots=True)
class SpecRowData:
    match_key: str
    fields: dict = field(default_factory=dict)
    raw: dict = field(default_factory=dict)


@dataclass(slots=True)
class SpecParseResult:
    rows: list[SpecRowData] = field(default_factory=list)
    headers: list[str] = field(default_factory=list)
    # Какая колонка распозналась как какое поле — показывается пользователю,
    # чтобы он видел, что именно система прочитала.
    column_map: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


def _norm_header(value: str) -> str:
    return re.sub(r"[^0-9a-zа-яё]+", " ", str(value).lower()).strip()


def build_column_map(headers: list[str]) -> dict[str, str]:
    """Заголовок из файла → поле детали."""
    aliases = spec_config().get("columns", {}) or {}
    mapping: dict[str, str] = {}
    for header in headers:
        norm = _norm_header(header)
        if not norm:
            continue
        for field_name, variants in aliases.items():
            if norm in {_norm_header(v) for v in variants}:
                mapping[header] = field_name
                break
    return mapping


def _coerce(field_name: str, value) -> object:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    if field_name in _NUMERIC_FIELDS:
        try:
            return float(text.replace(",", "."))
        except ValueError:
            return None
    if field_name in _INT_FIELDS:
        try:
            return int(float(text.replace(",", ".")))
        except ValueError:
            return None
    if field_name == "grain":
        return _grain_value(text)
    return text


def _grain_value(text: str) -> str:
    values = spec_config().get("grain_values", {}) or {}
    norm = _norm_header(text)
    for mode, variants in values.items():
        if norm in {_norm_header(v) for v in variants}:
            return {
                "length": GrainMode.ALONG_LENGTH,
                "width": GrainMode.ALONG_WIDTH,
                "none": GrainMode.NONE,
            }.get(mode, GrainMode.NONE)
    return GrainMode.NONE


def _match_key(fields: dict) -> str | None:
    for candidate in spec_config().get("match_by", []) or []:
        value = fields.get(candidate)
        if value:
            if candidate == "source_file":
                return PurePosixPath(str(value).replace("\\", "/")).stem.lower()
            return str(value).strip().lower()
    return None


def _rows_to_result(rows: list[dict], headers: list[str]) -> SpecParseResult:
    result = SpecParseResult(headers=headers)
    result.column_map = build_column_map(headers)
    if not result.column_map:
        result.warnings.append(
            "Ни одна колонка спецификации не распознана. "
            "Добавьте заголовки в config/spec_columns.yaml."
        )
        return result

    for index, raw in enumerate(rows, start=2):
        fields: dict = {}
        for header, field_name in result.column_map.items():
            fields[field_name] = _coerce(field_name, raw.get(header))
        key = _match_key(fields)
        if not key:
            result.warnings.append(
                f"Строка {index}: нечем связать с DXF "
                "(нет ни имени файла, ни артикула, ни наименования)."
            )
            continue
        result.rows.append(SpecRowData(match_key=key, fields=fields, raw=dict(raw)))
    return result


def parse_csv(data: bytes) -> SpecParseResult:
    cfg = spec_config()
    text = None
    for encoding in cfg.get("csv_encodings", ["utf-8"]):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        return SpecParseResult(warnings=["Не удалось определить кодировку CSV."])

    sample = text[:4096]
    delimiter = max(
        cfg.get("csv_delimiters", [";", ","]),
        key=lambda d: sample.count(d),
    )
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    rows = list(reader)
    headers = list(reader.fieldnames or [])
    return _rows_to_result(rows, headers)


def parse_xlsx(data: bytes) -> SpecParseResult:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return SpecParseResult(warnings=["Пустой файл спецификации."])

    headers = [str(h) if h is not None else "" for h in header_row]
    rows = [
        dict(zip(headers, values, strict=False))
        for values in rows_iter
        if any(v is not None and str(v).strip() for v in values)
    ]
    return _rows_to_result(rows, headers)


def parse_xml(data: bytes) -> SpecParseResult:
    cfg = spec_config().get("xml", {}) or {}
    try:
        root = ET.fromstring(data)
    except ET.ParseError as exc:
        return SpecParseResult(warnings=[f"Некорректный XML: {exc}"])

    elements: list[ET.Element] = []
    for xpath in str(cfg.get("row_xpath", ".//row")).split("|"):
        try:
            elements.extend(root.findall(xpath.strip()))
        except SyntaxError:
            continue
    if not elements:
        return SpecParseResult(warnings=["В XML не найдено строк спецификации."])

    rows: list[dict] = []
    headers: list[str] = []
    for element in elements:
        row: dict = dict(element.attrib) if cfg.get("attribute_first", True) else {}
        for child in element:
            if child.text and child.text.strip():
                row.setdefault(child.tag, child.text.strip())
        rows.append(row)
        for key in row:
            if key not in headers:
                headers.append(key)
    return _rows_to_result(rows, headers)


def parse_spec(filename: str, data: bytes) -> SpecParseResult:
    suffix = PurePosixPath(filename).suffix.lower()
    if suffix == ".csv":
        return parse_csv(data)
    if suffix in {".xlsx", ".xls"}:
        return parse_xlsx(data)
    if suffix == ".xml":
        return parse_xml(data)
    return SpecParseResult(warnings=[f"Формат спецификации {suffix} не поддерживается."])


def index_rows(rows: list[SpecRowData]) -> dict[str, SpecRowData]:
    return {row.match_key: row for row in rows}


def lookup(index: dict[str, SpecRowData], filename: str, code: str | None = None):
    """Ищет строку спецификации для DXF: по имени файла, затем по артикулу."""
    stem = PurePosixPath(filename.replace("\\", "/")).stem.lower()
    if stem in index:
        return index[stem]
    if code and code.lower() in index:
        return index[code.lower()]
    return None
